from io import BytesIO
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def generate_shipment_excel(shipment):
    """Генерация накладной в формате Excel."""
    wb = load_workbook("templ_xlsx/pok.xlsx")
    ws = wb.active
    ws.title = "Накладная"

    # --- стили ---
    #title_font = Font(name="Arial", size=14, bold=True)
    header_font = Font(name="Times New Roman", size=10, color="000000")
    #ell_font = Font(name="Arial", size=10)
    #header_fill = PatternFill(start_color="2D5F2D", end_color="2D5F2D", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    # --- таблица ---
    row = 7
    for i, pack in enumerate(shipment.booking.packs, 1):
        ws.cell(row=row, column=1).value = i
        ws.cell(row=row, column=1).font = header_font
        ws.cell(row=row, column=1).border = thin_border
        ws.cell(row=row, column=2).value = pack.number_pac
        ws.cell(row=row, column=2).font = header_font
        ws.cell(row=row, column=2).border = thin_border
        ws.cell(row=row, column=3).value = pack.grade
        ws.cell(row=row, column=3).font = header_font
        ws.cell(row=row, column=3).border = thin_border
        ws.cell(row=row, column=4).value = pack.width_nominal_for_pok
        ws.cell(row=row, column=4).font = header_font
        ws.cell(row=row, column=4).border = thin_border
        ws.cell(row=row, column=5).value = pack.width_fact_for_pok
        ws.cell(row=row, column=5).font = header_font
        ws.cell(row=row, column=5).border = thin_border
        ws.cell(row=row, column=6).value = pack.thickness_nominal_for_pok
        ws.cell(row=row, column=6).font = header_font
        ws.cell(row=row, column=6).border = thin_border
        ws.cell(row=row, column=7).value = pack.thickness_fact_for_pok
        ws.cell(row=row, column=7).font = header_font
        ws.cell(row=row, column=7).border = thin_border
        ws.cell(row=row, column=8).value = pack.length_nominal
        ws.cell(row=row, column=8).font = header_font
        ws.cell(row=row, column=8).border = thin_border
        ws.cell(row=row, column=9).value = pack.length_fact
        ws.cell(row=row, column=9).font = header_font
        ws.cell(row=row, column=9).border = thin_border
        ws.cell(row=row, column=10).value = pack.quantity_layers
        ws.cell(row=row, column=10).font = header_font
        ws.cell(row=row, column=10).border = thin_border
        ws.cell(row=row, column=11).value = pack.quantity_in_layer
        ws.cell(row=row, column=11).font = header_font
        ws.cell(row=row, column=11).border = thin_border
        ws.cell(row=row, column=12).value = "92%"
        ws.cell(row=row, column=12).font = Font(name="Arial", size=11, color="F80000")
        ws.cell(row=row, column=12).border = thin_border
        ws.cell(row=row, column=13).value = pack.volume_nominal_for_pok
        ws.cell(row=row, column=13).font = header_font
        ws.cell(row=row, column=13).border = thin_border
        ws.cell(row=row, column=14).value = pack.volume_fact_for_pok
        ws.cell(row=row, column=14).font = header_font
        ws.cell(row=row, column=14).border = thin_border
        ws.cell(row=row, column=15).value = pack.weight_nominal_for_pok
        ws.cell(row=row, column=15).font = header_font
        ws.cell(row=row, column=15).border = thin_border
        ws.cell(row=row, column=16).value = shipment.cont_number
        ws.cell(row=row, column=16).font = header_font
        ws.cell(row=row, column=16).border = thin_border
        ws.cell(row=row, column=17).value = shipment.shipped_at.strftime('%d.%m.%Y')
        ws.cell(row=row, column=17).font = header_font
        ws.cell(row=row, column=17).border = thin_border

        row += 1

    # --- итого ---
    # ws.cell(row=row, column=11, value="ИТОГО:").font = Font(name="Arial", size=11, bold=True)
    ws.cell(row=row, column=13, value=round(shipment.booking.total_volume_nominal_for_pok, 4)).font = Font(
        name="Arial", size=10, bold=True, color="F80000"
    )
    ws.cell(row=row, column=14, value=round(shipment.booking.total_volume_fact_for_pok, 4)).font = Font(
        name="Arial", size=8, bold=False , color="000000"
    )
    ws.cell(row=row, column=15, value=round(shipment.booking.total_weight_nominal_for_tam, 1)).font = Font(
        name="Arial", size=10, bold=True, color="F80000"
    )

    # --- ширины колонок ---
    # widths = [5, 14, 16, 8, 10, 10, 9, 9, 10, 12, 14]
    # for i, w in enumerate(widths, 1):
    #     ws.column_dimensions[chr(64 + i)].width = w

    # сохранение в буфер
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf